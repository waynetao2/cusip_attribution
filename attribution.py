"""
CUSIP-level asset attribution (rates vs spreads) - quarterly batch process.

Replaces the manual Cusip_level_rec_ParQuotes.xlsx build:
  1. Pull prior/current quarter positions, spread convexity and Axis cell mapping (SQL).
  2. Build monthly par curves (UST, JGB) by linear interpolation of par quotes.
  3. Per CUSIP x portfolio: rate impact, spread impact, actual MV change, residual.
  4. True-up rate/spread so they sum to actual (Bryan Goodwin's vBG proposal), with options.
  5. Data-quality flags, coverage check, summaries by model block x sector.
  6. Write one Excel output + CSV extract.

Usage
  python attribution.py --config config/26Q2.yaml                       # pull from database
  python attribution.py --config config/26Q2.yaml --from-workbook X.xlsx # rebuild from a prior workbook
  add --legacy to reproduce the delivered workbook's numbers exactly (for parallel runs).
"""
from __future__ import annotations

import argparse
import datetime as dt
import logging
from pathlib import Path

import numpy as np
import pandas as pd
import yaml

log = logging.getLogger("attribution")
HERE = Path(__file__).resolve().parent
KEY = ["riskid", "portfoliocode"]
NUM_COLS = ["WAL", "MV_CLEAN", "EFF_DUR", "EFF_CONV", "SPREAD_DUR", "OAS",
            "CurrentFace", "US_GAAP_Book_Value", "SPREAD_CONVEXITY"]


# --------------------------------------------------------------------------- data
def load_from_db(cfg: dict) -> dict[str, pd.DataFrame]:
    from sqlalchemy import create_engine, text
    eng = create_engine(cfg["db_url"])
    q = lambda f, d: pd.read_sql(text((HERE / "sql" / f).read_text()), eng, params={"asofdate": d})
    out = {}
    for tag, d in (("prior", cfg["prior_date"]), ("current", cfg["current_date"])):
        pos = q("01_positions.sql", d)
        if tag == "prior":
            sc = q("02_spread_convexity.sql", d)
            pos = pos.merge(sc.drop_duplicates(KEY), on=KEY, how="left")
        out[tag] = pos
    out["axis"] = q("03_axis_cell_map.sql", cfg["current_date"])
    try:
        out["igc_totals"] = q("04_igc_block_totals.sql", cfg["current_date"])
    except Exception as e:  # optional control totals
        log.warning("Coverage control totals not loaded: %s", e)
    return out


def load_from_workbook(path: str) -> dict[str, pd.DataFrame]:
    """Rebuild inputs from a delivered Cusip_level_rec_ParQuotes workbook (parallel-run / backfill)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = lambda ws, c0, c1: pd.DataFrame(
        [r[c0:c1] for r in wb[ws].iter_rows(min_row=2, values_only=True)])
    cols9 = ["riskid", "portfoliocode", "saa_level1", "WAL", "MV_CLEAN", "EFF_DUR", "EFF_CONV", "SPREAD_DUR", "OAS"]
    q1 = rows("Q1 data", 0, 9); q1.columns = cols9
    sc = rows("Q1 data", 10, 13); sc.columns = ["riskid", "portfoliocode", "SPREAD_CONVEXITY"]
    q2 = rows("Q2 data", 0, 9); q2.columns = cols9
    q1 = q1.merge(sc.dropna(subset=["riskid"]).drop_duplicates(KEY), on=KEY, how="left")
    if "par book" in wb.sheetnames:   # tab added in the vBG version
        pb = rows("par book", 0, 20)
        par2 = pb[[0, 1, 2]].set_axis(KEY + ["CurrentFace"], axis=1)
        par1 = pb[[5, 6, 7]].set_axis(KEY + ["CurrentFace"], axis=1)
        bv2 = pb[[12, 13, 14]].set_axis(KEY + ["US_GAAP_Book_Value"], axis=1)
        bv1 = pb[[17, 18, 19]].set_axis(KEY + ["US_GAAP_Book_Value"], axis=1)
        agg = lambda d: d.dropna(subset=["riskid"]).groupby(KEY, as_index=False).sum(numeric_only=True)
        q1 = q1.merge(agg(par1), on=KEY, how="left").merge(agg(bv1), on=KEY, how="left")
        q2 = q2.merge(agg(par2), on=KEY, how="left").merge(agg(bv2), on=KEY, how="left")
    cm = rows("Cell mapping script", 18, 21); cm.columns = ["cell_name", "subfund_name_input_bel", "Model"]
    cm = cm.dropna(subset=["cell_name"])
    cm["riskid"] = cm["cell_name"].str[19:28]
    cm["portfoliocode"] = "TR_" + cm["cell_name"].str[28:31]
    return {"prior": q1, "current": q2, "axis": cm}


def clean_positions(df: pd.DataFrame, scale: dict, legacy: bool = False) -> pd.DataFrame:
    df = df.replace("NULL", np.nan).copy()
    for c in NUM_COLS:
        if c in df:
            df[c] = pd.to_numeric(df[c], errors="coerce")
        else:
            df[c] = np.nan
    df = df[df["riskid"].notna()]
    # aggregate duplicate lots: MV/par/book additive, analytics MV-weighted
    w = df["MV_CLEAN"].fillna(0)
    add = ["MV_CLEAN", "CurrentFace", "US_GAAP_Book_Value"]
    wav = ["WAL", "EFF_DUR", "EFF_CONV", "SPREAD_DUR", "OAS", "SPREAD_CONVEXITY"]
    if df.duplicated(KEY).any():
        log.info("Aggregating %d duplicate lots", int(df.duplicated(KEY).sum()))
        tmp = df.assign(**{f"_w_{c}": df[c] * w for c in wav}, _w=w)
        g = tmp.groupby(KEY, as_index=False)
        out = g[add + [f"_w_{c}" for c in wav] + ["_w"]].sum(min_count=1)
        for c in wav:
            out[c] = np.where(out["_w"].abs() > 0, out[f"_w_{c}"] / out["_w"], np.nan)
        if legacy:   # workbook SUMIFS adds analytics across duplicate lots
            out[wav] = g[wav].sum(min_count=1)[wav].values
        out = out.merge(g["saa_level1"].first(), on=KEY)
        df = out[KEY + ["saa_level1"] + add + wav]
    f = df["portfoliocode"].map(scale).fillna(1.0)
    for c in add:
        df[c] = df[c] * f
    return df


# --------------------------------------------------------------------------- curves
def monthly_curve(tenors, rates, months=840) -> np.ndarray:
    """Value at t = 1..months months; linear between nodes, flat outside (matches the 'rates' tab)."""
    t = np.arange(1, months + 1) / 12.0
    return np.interp(t, np.asarray(tenors, float), np.asarray(rates, float))


def rate_shock_bp(wal: pd.Series, spec: dict, legacy: bool) -> pd.Series:
    d = (monthly_curve(spec["tenors"], spec["current"]) - monthly_curve(spec["tenors"], spec["prior"])) * 1e4
    # Excel VLOOKUP(approx) on a grid built by repeated +1/12 -> floor to the month
    grid = np.cumsum(np.full(len(d), 1 / 12.0))
    idx = np.searchsorted(grid, wal.fillna(-1).to_numpy(), side="right") - 1
    out = np.where(idx >= 0, d[np.clip(idx, 0, len(d) - 1)], np.nan if legacy else d[0])
    return pd.Series(out, index=wal.index)


# --------------------------------------------------------------------------- universe
def build_universe(p: pd.DataFrame, c: pd.DataFrame, remap: dict) -> pd.DataFrame:
    p = p.copy()
    p["portfolio_current"] = p["portfoliocode"].map(remap).fillna(p["portfoliocode"])
    m = p.merge(c, left_on=["riskid", "portfolio_current"], right_on=KEY, how="outer",
                suffixes=("_1", "_2"), indicator=True)
    m["riskid"] = m["riskid"].fillna(m.get("riskid_2"))
    m["portfoliocode"] = m["portfoliocode_1"].fillna(m["portfoliocode_2"])
    m["portfolio_current"] = m["portfolio_current"].fillna(m["portfoliocode_2"])
    m["status"] = np.select([m["_merge"] == "both", m["_merge"] == "left_only"],
                            ["Continuing", "Sold/Matured"], "New")
    m.loc[(m["status"] == "Continuing") & (m["portfoliocode_1"] != m["portfoliocode_2"]), "status"] = "Transferred"

    # auto-detect portfolio transfers (e.g. CFA -> CFM): same riskid, unique unmatched counterpart
    sold, new = m[m.status == "Sold/Matured"], m[m.status == "New"]
    s1 = sold.groupby("riskid").filter(lambda g: len(g) == 1)
    n1 = new.groupby("riskid").filter(lambda g: len(g) == 1)
    pairs = s1[["riskid"]].reset_index().merge(n1[["riskid"]].reset_index(), on="riskid")
    if len(pairs):
        log.info("Auto-matched %d portfolio transfers by riskid", len(pairs))
        cur_cols = [c for c in m.columns if c.endswith("_2")]
        for _, r in pairs.iterrows():
            m.loc[r.index_x, cur_cols] = m.loc[r.index_y, cur_cols].values
            m.loc[r.index_x, "portfolio_current"] = m.loc[r.index_y, "portfoliocode_2"]
            m.loc[r.index_x, "status"] = "Transferred"
        m = m.drop(index=pairs.index_y.values)
    m["saa_level1"] = m["saa_level1_1"].fillna(m["saa_level1_2"])
    return m.drop(columns=["_merge", "riskid_1", "riskid_2"], errors="ignore").reset_index(drop=True)


# --------------------------------------------------------------------------- calc
def taylor(mv, dur, conv, shock_bp):
    dy = shock_bp / 1e4
    return (-dur * dy + 0.5 * conv * dy ** 2) * mv


def calculate(u: pd.DataFrame, axis: pd.DataFrame, cfg: dict, legacy: bool) -> pd.DataFrame:
    dq = cfg["dq"]
    ax = axis.drop_duplicates(KEY)[KEY + ["Model"]].rename(columns={"portfoliocode": "portfolio_current"})
    u = u.merge(ax, on=["riskid", "portfolio_current"], how="left")
    u.loc[u["saa_level1"] == "Hedges", "Model"] = ""
    u["Model"] = u["Model"].fillna("").replace("", "Unmapped")

    f0 = lambda s: s.fillna(0.0)
    mv1, mv2 = f0(u["MV_CLEAN_1"]), f0(u["MV_CLEAN_2"])
    par1, par2 = f0(u["CurrentFace_1"]), f0(u["CurrentFace_2"])
    bv1, bv2 = f0(u["US_GAAP_Book_Value_1"]), f0(u["US_GAAP_Book_Value_2"])
    cont = u["status"].isin(["Continuing", "Transferred"]) & (mv2 != 0)

    # rate component
    shock = pd.Series(np.nan, index=u.index)
    for name, spec in cfg["curves"].items():
        sel = u["Model"].map(cfg.get("curve_by_model", {})).fillna("UST") == name
        shock[sel] = rate_shock_bp(u.loc[sel, "WAL_1"], spec, legacy)
    u["dRate_bp"] = np.where(cont, shock, 0.0)
    k = 1.0 if legacy else float(cfg.get("convexity_scale", 1.0))
    u["MV_rate"] = taylor(mv1, f0(u["EFF_DUR_1"]), k * f0(u["EFF_CONV_1"]), u["dRate_bp"]).fillna(0.0)

    # spread component
    oas1, oas2 = u["OAS_1"], u["OAS_2"]
    u["dOAS_bp"] = np.where(cont, (oas2 - oas1), 0.0)
    u["dOAS_bp"] = u["dOAS_bp"].fillna(0.0)
    u["MV_spread"] = taylor(mv1, f0(u["SPREAD_DUR_1"]), k * f0(u["SPREAD_CONVEXITY_1"]), u["dOAS_bp"]).fillna(0.0)

    # actual economic change
    u["dMV"] = mv2 - mv1
    u["dPar"] = par2 - par1
    u["dBook"] = bv2 - bv1
    legacy_act = u["dMV"] - u["dPar"]
    # non-USD blocks (JPI): par is local face, MV is USD -> strip FX so price change is in local terms
    fx1 = u["Model"].map({k: v["prior"] for k, v in (cfg.get("fx") or {}).items()}).fillna(1.0)
    fx2 = u["Model"].map({k: v["current"] for k, v in (cfg.get("fx") or {}).items()}).fillna(1.0)
    price_ok = (par1 > 0) & (par2 > 0)
    loc_px_chg = (mv2 * fx2) / par2.where(par2 != 0) - (mv1 * fx1) / par1.where(par1 != 0)
    price_act = np.where(price_ok, par1 * loc_px_chg / fx1, legacy_act)
    u["FX effect"] = np.where(cont & price_ok & (fx1 != fx2),
                              mv2 * par1 / par2.where(par2 != 0) - mv1 - price_act, 0.0)
    u["Actual"] = np.where(cont, legacy_act if cfg["actual_method"] == "legacy" else price_act, 0.0)
    u["Actual_legacy"] = np.where(cont, legacy_act, 0.0)
    u["Paydown/Trade effect"] = np.where(cont, u["dMV"] - u["Actual"] - u["FX effect"], u["dMV"])

    # time / pull-to-par: discount or premium amortised linearly over WAL (one quarter)
    if cfg.get("pull_to_par") and not legacy:
        ok = cont & (par1 > 0) & (u["WAL_1"] > 0.25)
        u["MV_time"] = np.where(ok, (par1 / fx1 - mv1) / (u["WAL_1"] * 4), 0.0)
    else:
        u["MV_time"] = 0.0

    # data quality
    flags = pd.Series("", index=u.index)
    add = lambda m, t: flags.where(~m, flags + t + "; ")
    bad_oas = cont & ((oas1.abs() > dq["max_abs_oas_bp"]) | (oas2.abs() > dq["max_abs_oas_bp"]))
    flags = add(bad_oas, "OAS outside +/-%d bp" % dq["max_abs_oas_bp"])
    flags = add(cont & (u["dOAS_bp"].abs() > dq["max_abs_doas_bp"]), "dOAS > %d bp" % dq["max_abs_doas_bp"])
    flags = add(cont & u["EFF_DUR_1"].isna(), "missing prior duration")
    flags = add(cont & u["WAL_1"].isna(), "missing WAL")
    flags = add(cont & u["SPREAD_DUR_1"].notna() & (u["SPREAD_DUR_1"] > 0) & u["SPREAD_CONVEXITY_1"].isna(),
                "missing spread convexity")
    flags = add(u["status"] == "Transferred", "portfolio transfer")
    mapped_ports = set(u.loc[u["Model"] != "Unmapped", "portfolio_current"])
    flags = add((u["Model"] == "Unmapped") & u["portfolio_current"].isin(mapped_ports) & (mv1 + mv2 != 0)
                & (u["saa_level1"] != "Hedges"), "no Axis cell in a modeled portfolio")

    # bad analytics: do not trust the spread estimate; spread becomes the plug
    u["Estimate_raw"] = u["MV_rate"] + u["MV_spread"]
    if not legacy:
        u.loc[bad_oas, "MV_spread"] = (u["Actual"] - u["MV_rate"] - u["MV_time"])[bad_oas]
    u["Estimate"] = u["MV_rate"] + u["MV_spread"] + u["MV_time"]
    u["Residual"] = u["Actual"] - u["Estimate"]
    big = cont & ((u["Residual"].abs() > dq["miss_flag_abs"]) |
                  (u["Residual"].abs() > dq["miss_flag_pct_mv"] * mv1.abs()))
    flags = add(big, "large residual")
    u["DQ_flags"] = flags.str.rstrip("; ")

    # true-up so components tie to actual
    r, s, res = u["MV_rate"], u["MV_spread"], u["Residual"]
    meth = cfg["true_up_method"]
    if meth == "abs_weight":
        wr, ws = r.abs(), s.abs()
    elif meth == "duration":
        wr = (f0(u["EFF_DUR_1"]) * u["dRate_bp"]).abs()
        ws = (f0(u["SPREAD_DUR_1"]) * u["dOAS_bp"]).abs()
    else:
        wr = ws = pd.Series(0.0, index=u.index)
    tot = wr + ws
    u["Rate_trued"] = np.where(tot > 0, r + wr / tot.where(tot > 0) * res, r)
    u["Spread_trued"] = np.where(tot > 0, s + ws / tot.where(tot > 0) * res, s)
    u["Other_unallocated"] = u["Actual"] - u["MV_time"] - u["Rate_trued"] - u["Spread_trued"]   # ties every row
    return u


# --------------------------------------------------------------------------- reporting
VALS = ["MV_CLEAN_1", "MV_CLEAN_2", "MV_rate", "MV_spread", "MV_time", "Estimate", "Actual", "Residual",
        "Rate_trued", "Spread_trued", "Other_unallocated", "FX effect", "Paydown/Trade effect", "dBook"]


def summarise(u: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    s = u.groupby(by, dropna=False)[VALS].sum().reset_index()
    s["Rate_bps_of_MV"] = s["Rate_trued"] / s["MV_CLEAN_1"].where(s["MV_CLEAN_1"] != 0) * 1e4
    s["Spread_bps_of_MV"] = s["Spread_trued"] / s["MV_CLEAN_1"].where(s["MV_CLEAN_1"] != 0) * 1e4
    s["Estimate_hit_ratio"] = s["Estimate"] / s["Actual"].where(s["Actual"] != 0)
    return s


def checks(u: pd.DataFrame, data: dict) -> pd.DataFrame:
    rows = [
        ("Rows in universe", len(u), ""),
        ("Continuing / Transferred / New / Sold", " / ".join(str((u.status == k).sum()) for k in
                                                           ["Continuing", "Transferred", "New", "Sold/Matured"]), ""),
        ("Rate + Spread + Time + Other - Actual (should be 0)", float((u.Rate_trued + u.Spread_trued + u.MV_time + u.Other_unallocated - u.Actual).abs().max()), "max abs, $"),
        ("Duplicate keys", int(u.duplicated(["riskid", "portfolio_current"]).sum()), ""),
        ("Rows with DQ flags", int((u.DQ_flags != "").sum()), ""),
        ("Unmapped MV (current)", float(u.loc[u.Model == "Unmapped", "MV_CLEAN_2"].sum()), "$"),
    ]
    chk = pd.DataFrame(rows, columns=["Check", "Value", "Unit"])
    if "igc_totals" in data:
        cov = summarise(u, ["Model"])[["Model", "MV_CLEAN_2"]].merge(
            data["igc_totals"].rename(columns={"model_block": "Model"}), on="Model", how="outer")
        cov["Coverage_ex_cash"] = cov["MV_CLEAN_2"] / (cov["MV_CLEAN"] - cov["CASH_MV"])
        chk = pd.concat([chk, cov.assign(Check="Coverage " + cov.Model.astype(str))[["Check", "Coverage_ex_cash"]]
                        .rename(columns={"Coverage_ex_cash": "Value"})], ignore_index=True)
    return chk


def write_outputs(u, cfg, data, outdir: Path, legacy: bool):
    outdir.mkdir(parents=True, exist_ok=True)
    tag = cfg["quarter_label"] + ("_legacy" if legacy else "")
    xl = outdir / f"Cusip_level_attribution_{tag}.xlsx"
    exhibit_cols = ["riskid", "portfoliocode", "portfolio_current", "status", "Model", "saa_level1", "WAL_1",
                    "OAS_1", "OAS_2", "MV_CLEAN_1", "MV_CLEAN_2", "CurrentFace_1", "CurrentFace_2",
                    "US_GAAP_Book_Value_1", "US_GAAP_Book_Value_2", "EFF_DUR_1", "EFF_CONV_1", "dRate_bp",
                    "MV_rate", "SPREAD_DUR_1", "SPREAD_CONVEXITY_1", "dOAS_bp", "MV_spread", "MV_time", "Estimate",
                    "Actual", "Actual_legacy", "Residual", "Rate_trued", "Spread_trued", "Other_unallocated",
                    "FX effect", "Paydown/Trade effect", "DQ_flags"]
    with pd.ExcelWriter(xl, engine="xlsxwriter") as w:
        info = pd.DataFrame({"Item": ["Quarter", "Prior date", "Current date", "True-up method", "Actual method",
                                      "Legacy mode", "Run time"],
                             "Value": [cfg["quarter_label"], cfg["prior_date"], cfg["current_date"],
                                       cfg["true_up_method"], cfg["actual_method"], legacy,
                                       dt.datetime.now().strftime("%Y-%m-%d %H:%M")]})
        info.to_excel(w, sheet_name="Run info", index=False)
        summarise(u, ["Model"]).to_excel(w, sheet_name="Summary by block", index=False)
        summarise(u, ["Model", "saa_level1"]).to_excel(w, sheet_name="Block x sector", index=False)
        checks(u, data).to_excel(w, sheet_name="Checks", index=False)
        u.loc[u.DQ_flags != "", exhibit_cols].sort_values("Residual", key=abs, ascending=False) \
            .to_excel(w, sheet_name="DQ exceptions", index=False)
        u[exhibit_cols].to_excel(w, sheet_name="exhibit", index=False)
        for sh in w.sheets.values():
            sh.freeze_panes(1, 0)
            sh.set_column(0, 40, 16)
    u[exhibit_cols].to_csv(outdir / f"cusip_attribution_{tag}.csv", index=False)
    log.info("Wrote %s", xl)
    return xl


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--from-workbook")
    ap.add_argument("--legacy", action="store_true", help="reproduce delivered workbook logic exactly")
    ap.add_argument("--outdir", default=str(HERE / "outputs"))
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    cfg = yaml.safe_load(Path(a.config).read_text())
    if a.legacy:
        cfg = {**cfg, "actual_method": "legacy", "true_up_method": "abs_weight"}
    data = load_from_workbook(a.from_workbook) if a.from_workbook else load_from_db(cfg)
    scale = cfg.get("portfolio_scale", {})
    p, c = clean_positions(data["prior"], scale, a.legacy), clean_positions(data["current"], scale, a.legacy)
    u = build_universe(p, c, cfg.get("portfolio_remap", {}) or {})
    if a.legacy:   # delivered workbook = inner join on riskid + portfolio, no transfer matching
        u = u[u.status == "Continuing"].reset_index(drop=True)
    u = calculate(u, data["axis"], cfg, a.legacy)
    return write_outputs(u, cfg, data, Path(a.outdir), a.legacy), u


if __name__ == "__main__":
    main()
